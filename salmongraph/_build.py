"""Offline, transactional dataset builder behind executor.build_dataset.

This creates a derived private build, not a publication or scientific review.
Source parsing follows the saved spec; the public SDK's inferred semantics are
never adopted. Every exported file except the manifest hashes into the manifest.
"""

from __future__ import annotations

from contextlib import contextmanager
from collections.abc import Mapping
import csv
import datetime as dt
from decimal import Decimal, InvalidOperation
import hashlib
import importlib.util
import inspect
import io
import json
import os
from pathlib import Path
import re
import shutil
import socket
import sys
import tempfile
import warnings

import pandas as pd

from . import __version__, config, iri, review
from .graph import GraphBuilder, expand
from .io import det
from .spec import SpecError, load_spec, spec_sha256, validate_spec
from .vocab import VocabIndex

_FORMAT = "salmongraph-build/1"


def build_dataset(spec, *, build_date: str, vocab=None, raw_dir=None, out_dir=None,
                  run=None, validate: bool = True):
    """Execute the public executor contract using genuine local SDK/gate calls."""
    from .executor import BUILD_FILES, BuildResult, build_dir, write_observations
    from ._compile import compile_dataset_graph
    from . import validate as validation
    import metasalmonpy as ms

    _build_date(build_date)
    if not isinstance(validate, bool):
        raise SpecError("build: validate must be True or False")
    vocab = vocab if vocab is not None else VocabIndex(offline=True)
    with config.override(vocab.cfg), _offline() as no_network, _vendored_schema():
        saved = _saved_spec(spec)
        problems = validate_spec(spec, vocab)
        if problems:
            raise SpecError("build: invalid spec:\n" + "\n".join(problems))
        raw = _safe_path(raw_dir or config.ROOT / "data/raw" / spec.slug, existing=True)
        if not raw.is_dir():
            raise SpecError("build: raw_dir must be a local directory")
        destination = _safe_path(out_dir or build_dir(spec.slug, spec.version))
        _separate(destination, raw, saved.path)
        decisions = _decision_state()
        sources = _raw_records(spec, raw)
        parser = _parser_source(spec)
        if parser:
            _separate(destination, parser[0])
        sdk = _sdk_snapshot(ms)
        implementation = _implementation_snapshot(validate)
        prior = _prior_build(destination, spec)
        frames = _frames(spec, raw, parser)
        no_network()
        destination.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=".salmongraph-build-", dir=destination.parent) as temporary:
            stage = Path(temporary)
            sdp = stage / "sdp"
            sdk_result = _package(spec, frames, sdp, build_date)
            builder = GraphBuilder(vocab, cfg=vocab.cfg)
            graphs = compile_dataset_graph(spec, sdp, builder=builder, build_date=build_date, run=run)
            builder.assert_vocab()
            shacl = validation.assert_valid(builder.quads) if validate else None
            count = write_observations(spec, frames, stage / "observations.parquet")
            builder.write_trig(stage / "dataset.trig")
            queue = review.review_queue(builder.quads)
            review.write_review_queue(queue, stage / "review-queue.csv")
            counts = review.status_counts(builder.quads)
            report = _report(spec, frames, builder, sdk_result, count, counts, validate, shacl)
            det.write_json(stage / "report.json", report)
            det.write_text(stage / "report.md", _report_markdown(report))
            manifest = {
                "schema": _FORMAT, "slug": spec.slug, "version": spec.version,
                "build_date": build_date, "dataset_iri": iri.dataset(spec.slug, spec.version),
                "configuration": {"base": vocab.cfg.base, "profile_base": vocab.cfg.profile_base, "sdpv": vocab.cfg.sdpv},
                "manifest_hash": "excluded: manifest.json cannot hash itself",
                "spec": {"sha256": spec_sha256(spec)}, "sources": sources,
                "parser": {"kind": spec.parser, "source_sha256": parser[1] if parser else None},
                "vocabularies": vocab.pins_dict(), "tools": _tools(ms, sdk, implementation, validate),
                "sdp_schema": {"source": "vendored", "version": sdk_result["schema_version"]},
                "review_decisions": decisions, "quad_counts": builder.graphs(),
                "observations": count, "review_counts": counts,
                "validation": report["checks"], "publication_ready": report["publication_ready"],
                "publication_gaps": report["publication_gaps"],
                "files": _file_records(stage),
            }
            det.write_json(stage / "manifest.json", manifest)
            # Recheck all external mutable inputs immediately before publication.
            _saved_spec(spec)
            if _raw_records(spec, raw) != sources or _decision_state() != decisions:
                raise SpecError("build: source files or review decisions changed during build")
            if parser and det.sha256_file(parser[0]) != parser[1]:
                raise SpecError("build: custom parser changed during build")
            if _sdk_snapshot(ms) != sdk:
                raise SpecError("build: runtime SDK source changed during build")
            if _implementation_snapshot(validate) != implementation:
                raise SpecError("build: salmongraph implementation/schema/shapes changed during build")
            no_network()
            _publish(stage, destination, spec, prior)
        return BuildResult(spec.slug, spec.version, destination,
                           files={name: destination / name for name in BUILD_FILES}, graphs=graphs,
                           quads=len(builder), observations=count, review_counts=counts,
                           undeclared=[], shacl_conforms=shacl.conforms if shacl is not None else None)


def _build_date(value):
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        raise SpecError("build: build_date must be an explicit YYYY-MM-DD")
    try:
        dt.date.fromisoformat(value)
    except ValueError as exc:
        raise SpecError("build: invalid build_date") from exc


def _safe_path(path, *, existing=False):
    path = Path(path)
    if ".." in path.parts:
        raise SpecError("build: parent traversal is unsupported")
    path = path.absolute()
    for candidate in (path, *path.parents):
        if candidate.is_symlink():
            raise SpecError("build: symlink paths are unsupported")
    if existing and not path.exists():
        raise SpecError("build: a declared local input is missing")
    return path


def _separate(destination, *inputs):
    for source in (config.ROOT, *inputs):
        source = Path(source).absolute()
        if source == destination or source.is_relative_to(destination):
            raise SpecError("build: destination overlaps a source or project directory")
        if source != config.ROOT and source.is_dir() and destination.is_relative_to(source):
            raise SpecError("build: output cannot be inside raw input")


def _saved_spec(spec):
    digest = spec_sha256(spec)
    _safe_path(spec.path, existing=True)
    saved = load_spec(spec.path)
    if saved != spec or spec_sha256(saved) != digest:
        raise SpecError("build: in-memory spec differs from saved YAML; save and reload it")
    return saved


def _raw_records(spec, raw):
    records = []
    for item in sorted(spec.source["files"], key=lambda item: item["name"]):
        path = _safe_path(raw / item["name"], existing=True)
        if not path.is_file() or not path.is_relative_to(raw):
            raise SpecError("build: declared raw input must be a contained regular file")
        digest = det.sha256_file(path)
        if digest != item["sha256"]:
            raise SpecError(f"build: source SHA-256 mismatch for {item['name']!r}")
        records.append({"name": item["name"], "url": item["url"], "bytes": path.stat().st_size, "sha256": digest})
    return records


def _decision_state():
    path = _safe_path(config.ROOT / "reviews/decisions.csv")
    if not path.exists():
        return {"status": "none_present"}
    if not path.is_file():
        raise SpecError("build: review decisions path must be a regular file")
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = [row for row in csv.reader(handle) if any(value for value in row)]
    if rows and rows[0] != list(review.DECISION_COLUMNS):
        raise SpecError("build: unrecognized review decisions file; human-decision application is not implemented")
    if len(rows) > 1:
        raise SpecError("build: human review decisions exist but applying them is not implemented")
    return {"status": "no_decision_rows", "sha256": det.sha256_file(path)}


@contextmanager
def _offline():
    # Some libraries swallow connection errors. Count attempts as well as
    # refusing them, so no request/model fetch can silently become a pass.
    calls, originals = [], {}
    def blocked(*args, **kwargs):
        calls.append(True)
        raise SpecError("build: network access is forbidden; provide pinned local inputs")
    def check():
        if calls:
            raise SpecError("build: a component attempted forbidden network access")
    targets = [(socket, "create_connection"), (socket, "getaddrinfo"),
               (socket.socket, "connect"), (socket.socket, "connect_ex")]
    try:
        for owner, name in targets:
            originals[owner, name] = getattr(owner, name)
            setattr(owner, name, blocked)
        yield check
        check()
    finally:
        for (owner, name), original in originals.items():
            setattr(owner, name, original)


@contextmanager
def _vendored_schema():
    from metasalmonpy import sdp_schema
    # The public setter is used for both changes. Read the previous override
    # to restore None (environment-driven) exactly, not freeze its effective value.
    previous = getattr(sdp_schema, "_SOURCE_OVERRIDE", None)
    sdp_schema.set_sdp_schema_source("vendored")
    try:
        yield
    finally:
        sdp_schema.set_sdp_schema_source(previous)


def _parser_source(spec):
    if spec.parser in ("csv", "xlsx"):
        return None
    module_name, function = spec.parser.split(":", 1)
    parts = module_name.split(".")
    candidates = [config.ROOT.joinpath(*parts).with_suffix(".py"),
                  config.ROOT.joinpath(*parts, "__init__.py")]
    candidates = [path for path in candidates if path.exists() or path.is_symlink()]
    if len(candidates) != 1:
        raise SpecError("build: custom parser must resolve to one actual project-local module")
    path = _safe_path(candidates[0], existing=True)
    if not path.is_file() or not path.is_relative_to(config.ROOT):
        raise SpecError("build: custom parser is not project-local source")
    return path, det.sha256_file(path)


def _frames(spec, raw, parser):
    if parser:
        path, digest = parser
        name = "_salmongraph_parser_" + digest
        module_spec = importlib.util.spec_from_file_location(name, path)
        module = importlib.util.module_from_spec(module_spec)
        package = spec.parser.split(":", 1)[0].rpartition(".")[0]
        module.__package__ = package
        # Register the verified source while it executes (dataclasses and
        # ordinary Python module introspection need this), then restore the
        # caller's import state. No installed-module entrypoint is imported.
        previous, search_path = sys.modules.get(name), list(sys.path)
        sys.modules[name] = module
        sys.path.insert(0, str(config.ROOT))
        try:
            module_spec.loader.exec_module(module)
            function = getattr(module, spec.parser.split(":", 1)[1], None)
            if not inspect.isfunction(function) or Path(inspect.getsourcefile(function) or "") != path:
                raise SpecError("build: parser function must be defined in the verified project-local source")
            frames = function(spec=spec, raw_dir=raw)
        finally:
            sys.path[:] = search_path
            if previous is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = previous
        if not isinstance(frames, Mapping) or set(frames) != {table["id"] for table in spec.tables}:
            raise SpecError("build: custom parser must return exactly the declared table frames")
    else:
        frames = {}
        for table in spec.tables:
            path = raw / table["file"]
            if spec.parser == "csv":
                with path.open(newline="", encoding="utf-8-sig") as handle:
                    records = list(csv.reader(handle, strict=True))
            else:
                import openpyxl
                with path.open("rb") as handle:
                    workbook = openpyxl.load_workbook(handle, read_only=True, data_only=False)
                    try:
                        if table["sheet"] not in workbook.sheetnames:
                            raise SpecError("build: declared XLSX sheet is missing")
                        records = []
                        for row in workbook[table["sheet"]].iter_rows():
                            if any(cell.data_type == "f" for cell in row):
                                raise SpecError("build: XLSX formulas require an explicit reviewed parser")
                            records.append([cell.value for cell in row])
                    finally:
                        workbook.close()
            if not records or len(records[0]) != len(set(records[0])):
                raise SpecError("build: empty input or duplicate source headers")
            headers = records[0]
            if any(len(row) != len(headers) for row in records[1:]):
                raise SpecError("build: inconsistent source row width")
            frames[table["id"]] = pd.DataFrame(records[1:], columns=headers, dtype=object)
    output = {}
    for table in spec.tables:
        frame = frames[table["id"]]
        names = [column["name"] for column in table["columns"]]
        if not isinstance(frame, pd.DataFrame) or not frame.columns.is_unique or set(frame.columns) != set(names):
            raise SpecError("build: source frame columns must exactly match the declared columns")
        typed = pd.DataFrame({column["name"]: pd.Series(
            [_typed(value, column["datatype"], column["name"]) for value in frame[column["name"]]], dtype=object)
            for column in table["columns"]})
        key_values = []
        for row in typed[table["keys"]].itertuples(index=False, name=None):
            if any(value is None or value == "" for value in row):
                raise SpecError("build: source primary key is missing")
            key_values.append(tuple(iri.segment(value) for value in row))
        if len(key_values) != len(set(key_values)):
            raise SpecError("build: duplicate source primary key")
        output[table["id"]] = typed
    return output


def _typed(value, datatype, name):
    from .executor import _obs_missing, _obs_scalar
    value = _obs_scalar(value)
    if _obs_missing(value):
        return "" if datatype == "string" else None
    if datatype == "string":
        if not isinstance(value, str):
            raise SpecError(f"build: string source column {name!r} requires text; no identifier guessing")
        return value
    if datatype == "boolean":
        if isinstance(value, bool):
            return value
        if isinstance(value, str) and value.casefold() in ("true", "false"):
            return value.casefold() == "true"
        raise SpecError(f"build: invalid explicit boolean in {name!r}")
    if datatype == "date":
        if isinstance(value, dt.datetime):
            if value.time() != dt.time():
                raise SpecError(f"build: date column {name!r} would lose a time")
            return value.date()
        if isinstance(value, dt.date):
            return value
        if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise SpecError(f"build: date column {name!r} requires ISO YYYY-MM-DD")
        try:
            return dt.date.fromisoformat(value)
        except ValueError as exc:
            raise SpecError(f"build: invalid date in {name!r}") from exc
    if isinstance(value, bool):
        raise SpecError(f"build: boolean is not a numeric value in {name!r}")
    try:
        number = Decimal(str(value))
    except InvalidOperation as exc:
        raise SpecError(f"build: invalid declared numeric value in {name!r}") from exc
    if not number.is_finite():
        raise SpecError(f"build: non-finite declared number in {name!r}")
    if datatype == "integer":
        if number != int(number) or not -(2**63) <= number < 2**63:
            raise SpecError(f"build: integer is fractional or outside int64 in {name!r}")
        return int(number)
    converted = float(number)
    if not converted == converted or converted in (float("inf"), -float("inf")) or Decimal(str(converted)) != number:
        raise SpecError(f"build: numeric precision loss in {name!r}")
    return converted


def _package(spec, frames, path, build_date):
    import metasalmonpy as ms
    from metasalmonpy import sdp_schema
    dataset = pd.DataFrame([{"dataset_id": spec.slug, "title": spec.title, "description": spec.description,
                             "license": spec.source["licence"], "created": build_date,
                             "source_citation": spec.source["attribution"]}])
    tables, dictionary, codes = [], [], []
    for table in spec.tables:
        tables.append({"dataset_id": spec.slug, "table_id": table["id"], "file_name": table["id"] + ".csv",
                       "primary_key": ",".join(table["keys"])})
        for column in table["columns"]:
            decomposition = column.get("decomposition") or {}
            dictionary.append({"dataset_id": spec.slug, "table_id": table["id"], "column_name": column["name"],
                               "column_label": column.get("label", column["name"]), "column_description": column.get("description"),
                               "column_role": {"measure": "measurement", "dimension": "attribute"}.get(column["role"], column["role"]),
                               "value_type": column["datatype"], "required": column["name"] in table["keys"],
                               "term_iri": column.get("concept"), "unit_iri": column.get("unit"),
                               "property_iri": decomposition.get("property"), "entity_iri": decomposition.get("object_of_interest"),
                               "constraint_iri": ";".join(decomposition.get("constraints", [])) or None,
                               "statistical_modifier_iri": decomposition.get("statistical_modifier")})
            for value, target in column.get("codes", {}).items():
                codes.append({"dataset_id": spec.slug, "table_id": table["id"], "column_name": column["name"],
                              "code_value": value, "code_label": value, "term_iri": target})
    table_meta, dictionary = pd.DataFrame(tables), pd.DataFrame(dictionary)
    code_frame = pd.DataFrame(codes) if codes else None
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        # Exercise the public inference boundary without accepting its guessed
        # roles, semantic seeds or legacy crosswalk prefills as source facts.
        artifacts = ms.infer_salmon_datapackage_artifacts(
            frames, dataset_id=spec.slug, guess_types=False, seed_semantics=False,
            llm_assess=False, seed_verbose=False, semantic_code_scope="none",
            seed_dataset_meta=dataset, seed_table_meta=table_meta, seed_codes=code_frame)
        ms.write_salmon_datapackage(frames, dataset, table_meta, dictionary, code_frame, path=str(path))
        _extensions(ms, spec, path)
        checked = ms.validate_salmon_datapackage(path, require_iris=False)
        roundtrip = ms.read_salmon_datapackage(path)["resources"]
        _same_frames(frames, roundtrip)
        strict = []
        try:
            ms.validate_salmon_datapackage(path, require_iris=True)
        except ValueError as exc:
            strict = [str(exc).replace(str(path), "sdp")]
    issues = checked["semantic_validation"].get("issues")
    messages = sorted(set(str(value) for value in issues["message"])) if isinstance(issues, pd.DataFrame) and "message" in issues else []
    return {"structural_valid": True, "semantic_issues": messages, "strict_publication_issues": strict,
            "schema_version": sdp_schema.load_sdp_schema(quiet=True)["version"],
            "inference": {"seed_semantics": False, "llm_assess": False, "adopted_guesses": False}}


def _same_frames(expected, actual):
    if set(expected) != set(actual):
        raise SpecError("build: SDK round trip changed declared table keys")
    for table, frame in expected.items():
        other = actual[table]
        if list(frame.columns) != list(other.columns) or frame.shape != other.shape:
            raise SpecError("build: SDK round trip changed source table shape")
        for left, right in zip(frame.itertuples(index=False, name=None), other.itertuples(index=False, name=None)):
            for a, b in zip(left, right):
                from .executor import _obs_missing
                if _obs_missing(a) and _obs_missing(b):
                    continue
                if isinstance(a, str):
                    same = isinstance(b, str) and a == b
                else:
                    same = not _obs_missing(b) and a == b
                if not same:
                    raise SpecError("build: SDK round trip changed a source value; use an explicit lossless parser/SDK")


def _extensions(ms, spec, path):
    structures, components, decompositions = [], [], []
    tables = {table["id"]: table for table in spec.tables}
    for block in spec.observations:
        for measure in block["measures"]:
            structure = "measure_" + iri.stable_key(block["table"], *block["key"], measure)
            common = {"dataset_id": spec.slug, "table_id": block["table"], "observation_structure_id": structure}
            structures.append({**common, "structure_label": measure,
                               "structure_description": "Measure-specific structure declared in the saved dataset specification."})
            required = set(block["dimensions"]) & set(block["key"])
            required.add(block["year_column"])
            bindings = [(name, "dimension", True) for name in block["dimensions"] if name in required]
            bindings += [(name, "attribute", False) for name in block["dimensions"] if name not in required]
            bindings += [(measure, "measure", True)]
            for order, (name, role, needed) in enumerate(bindings, 1):
                components.append({**common, "component_order": order, "column_name": name,
                                   "component_role": role, "component_relation_iri": "", "required_when_observed": needed})
    if structures:
        ms.write_sdp_observation_structures(path, pd.DataFrame(structures), pd.DataFrame(components))
    for table in spec.tables:
        for column in table["columns"]:
            definition = column.get("decomposition")
            if not definition:
                continue
            parts = [("property", definition["property"]), ("entity", definition["object_of_interest"])]
            parts += [("constraint", value) for value in definition.get("constraints", [])]
            if definition.get("statistical_modifier"):
                parts.append(("statistical_modifier", definition["statistical_modifier"]))
            if column.get("unit"):
                parts.append(("unit", column["unit"]))
            for order, (role, target) in enumerate(parts, 1):
                decompositions.append({"dataset_id": spec.slug, "table_id": table["id"], "column_name": column["name"],
                    "measurement_concept_iri": column["concept"], "component_order": order, "component_role": role,
                    "component_status": "matched", "component_relation": "", "related_component_order": "",
                    "component_iri": target, "component_label": "", "rationale": "Explicit saved spec component; no review acceptance inferred.",
                    "source": spec.source["publisher"], "source_version": spec.version,
                    "source_url": spec.source["landing_page"], "provenance": "Saved spec sha256:" + spec_sha256(spec)})
    if decompositions:
        ms.write_sdp_measurement_decompositions(path, pd.DataFrame(decompositions))


def _report(spec, frames, builder, sdk, count, counts, validate, shacl):
    gaps = ["Source creator and inquiry contact name/email are not provided by this spec contract."]
    if spec.source.get("licence") is None:
        gaps.append("Source licence URI is unknown; no grant or public-domain status is inferred.")
    if spec.source["redistribution"] != "confirmed":
        gaps.append("Source redistribution is unconfirmed.")
    if counts.get("awaitingReview"):
        gaps.append(f"{counts['awaitingReview']} semantic targets await human review; no decisions were applied.")
    if sdk["semantic_issues"] or sdk["strict_publication_issues"]:
        gaps.append("SDP publication metadata/semantic validation has unresolved issues; see sdp checks.")
    for table in spec.tables:
        for column in table["columns"]:
            if column.get("codes"):
                gaps.append(f"{table['id']}.{column['name']}: code targets lack an explicit mapping predicate; no equivalence is asserted.")
    if not validate:
        gaps.append("SHACL was explicitly skipped; this build is not structurally certified.")
    return {"slug": spec.slug, "version": spec.version, "artifact_kind": "derived private package/release build",
            "source_dates": {"publication": None, "retrieval": None},
            "tables": [{"id": table["id"], "source_file": table["file"], "rows": len(frames[table["id"]])} for table in spec.tables],
            "observations": count, "quads": len(builder), "quad_counts": builder.graphs(), "review_counts": counts,
            "checks": {"raw_sha256": "verified_before_parse_and_before_publish", "saved_spec": "verified_before_parse_and_before_publish",
                       "sdp": sdk, "vocabulary": {"undeclared": []},
                       "shacl": {"status": "passed" if validate else "skipped", "conforms": shacl.conforms if shacl else None},
                       "cross_source": {"status": "not_run", "reason": "combine is not implemented", "expectations": list(spec.cross_source_checks)}},
            "publication_ready": False, "publication_gaps": sorted(set(gaps)),
            "scientific_review": "not performed; declared semantics and structural validation are not scientific approval"}


def _report_markdown(report):
    return (f"# {report['slug']} {report['version']}\n\nDerived private build.\n\n"
            f"Observations: {report['observations']}. RDF quads: {report['quads']}.\n\n"
            f"SHACL: {report['checks']['shacl']['status']}. Cross-source checks: not run.\n\n"
            "Publication gaps:\n\n" + "".join(f"- {gap}\n" for gap in report["publication_gaps"])
            + "\nNo scientific approval or human decision is implied.\n")


def _sdk_snapshot(ms):
    root = _safe_path(Path(ms.__file__).parent, existing=True)
    files = []
    for path in sorted(root.rglob("*")):
        relative = path.relative_to(root)
        if any(part.startswith(".") or part in ("tests", "docs", "__pycache__") for part in relative.parts):
            continue
        if path.is_symlink():
            raise SpecError("build: runtime SDK source symlinks are unsupported")
        if path.is_file() and (path.suffix == ".py" or relative.parts[0] == "data"):
            files.append({"path": relative.as_posix(), "sha256": det.sha256_file(path)})
    if not files:
        raise SpecError("build: actual runtime SDK source cannot be fingerprinted")
    return {"version": ms.__version__, "source_sha256": det.sha256_bytes(det.json_bytes(files)), "source_files": files}


def _implementation_snapshot(validated):
    root = _safe_path(Path(__file__).parent, existing=True)
    paths = list(root.rglob("*.py")) + [root / "spec.schema.json"]
    if validated:
        from .shapes import shape_files
        paths += shape_files()
    records = []
    for path in sorted(set(paths)):
        path = _safe_path(path, existing=True)
        if not path.is_relative_to(root) or not path.is_file():
            raise SpecError("build: implementation source is not a regular local package file")
        records.append({"path": path.relative_to(root).as_posix(), "sha256": det.sha256_file(path)})
    return {"version": __version__, "source_sha256": det.sha256_bytes(det.json_bytes(records)), "source_files": records}


def _tools(ms, sdk, implementation, validated):
    import pyarrow, pyoxigraph
    tools = {"salmongraph": implementation, "python": sys.version.split()[0], "pandas": pd.__version__,
             "pyarrow": pyarrow.__version__, "pyoxigraph": pyoxigraph.__version__, "metasalmonpy": sdk}
    if validated:
        import pyshacl, rdflib
        tools.update(pyshacl=pyshacl.__version__, rdflib=rdflib.__version__)
    return tools


def _file_records(root):
    records = []
    for path in sorted(root.rglob("*")):
        if path.is_symlink() or not (path.is_file() or path.is_dir()):
            raise SpecError("build: symlink/special entry in build tree")
        relative = path.relative_to(root).as_posix()
        if path.is_file() and relative != "manifest.json":
            records.append({"path": relative, "bytes": path.stat().st_size, "sha256": det.sha256_file(path)})
    return records


def _prior_build(path, spec):
    if not path.exists():
        return None
    if not path.is_dir() or (path / "manifest.json").is_symlink():
        raise SpecError("build: destination is not a recognized prior build")
    try:
        manifest_bytes = (path / "manifest.json").read_bytes()
        manifest = json.loads(manifest_bytes)
    except (OSError, ValueError) as exc:
        raise SpecError("build: existing destination has no recognized build manifest") from exc
    if manifest.get("schema") != _FORMAT or manifest.get("slug") != spec.slug or manifest.get("version") != spec.version:
        raise SpecError("build: destination belongs to a different/unrecognized build")
    declared = manifest.get("files")
    if (not isinstance(declared, list) or not all(isinstance(item, dict) and isinstance(item.get("path"), str)
                                                for item in declared)
            or len({item["path"] for item in declared}) != len(declared)):
        raise SpecError("build: invalid prior build file inventory")
    for item in declared:
        relative = Path(item.get("path", ""))
        if relative.is_absolute() or not relative.parts or ".." in relative.parts or relative.as_posix() == "manifest.json":
            raise SpecError("build: unsafe prior build inventory path")
    expected_dirs = {part.as_posix() for item in declared for part in Path(item["path"]).parents if part != Path(".")}
    actual_dirs = set()
    for candidate in path.rglob("*"):
        if candidate.is_symlink():
            raise SpecError("build: symlink in prior build")
        if candidate.is_dir():
            actual_dirs.add(candidate.relative_to(path).as_posix())
    if actual_dirs != expected_dirs or _file_records(path) != declared:
        raise SpecError("build: prior output is dirty or contains unrecognized files/directories")
    return det.sha256_bytes(manifest_bytes)


def _publish(stage, destination, spec, prior):
    if _prior_build(destination, spec) != prior:
        raise SpecError("build: destination changed during build")
    if prior is None:
        os.replace(stage, destination)
        return
    # Only a verified prior build is moved/deleted. Never remove an arbitrary
    # caller directory, and roll back if installing the new staged tree fails.
    backup = Path(tempfile.mkdtemp(prefix=".salmongraph-prior-", dir=destination.parent))
    backup.rmdir()  # Our own newly created empty reservation, not user content.
    os.replace(destination, backup)
    try:
        if _prior_build(backup, spec) != prior:
            raise SpecError("build: prior output changed during replacement")
        os.replace(stage, destination)
    except BaseException:
        os.replace(backup, destination)
        raise
    # Recheck before deleting the recognized prior build; preserve/roll back
    # anything that was edited concurrently rather than silently dropping it.
    try:
        if _prior_build(backup, spec) != prior:
            raise SpecError("build: prior output changed before cleanup")
    except BaseException:
        os.replace(destination, stage)
        os.replace(backup, destination)
        raise
    shutil.rmtree(backup)
